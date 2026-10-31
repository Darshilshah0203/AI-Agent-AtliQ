"""Half-page Word escalation report generator (.docx)."""

import os
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from src.alerts.resolver import ConsolidatedAlert, AlertType
from src.ai.drop_investigator import DropInvestigationResult
from src.reports.formatting import (
    RGB_RED,
    RGB_BLUE,
    RGB_GRAY,
    RGB_BLACK,
    HEX_RED,
    HEX_LIGHT_BG,
    HEX_BORDER,
)


def sanitize_filename_part(part: str) -> str:
    """Sanitize string for Windows filesystem safety."""
    cleaned = re.sub(r'[\\/*?:"<>|]', "", part)
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned.strip("._")


def get_escalation_docx_path(alert: ConsolidatedAlert, base_dir: str = "escalations") -> str:
    """Generate sanitized output path for escalation .docx file.
    
    Naming rules (PRD 8.3):
    - Combined / Market-Channel: escalations/<market>_<channel>_<date>.docx
    - Standalone customer: escalations/<market>_<channel>_<customer>_<date>.docx
    """
    market_clean = sanitize_filename_part(alert.market)
    channel_clean = sanitize_filename_part(alert.channel)
    date_str = alert.date.isoformat()

    if alert.alert_type == AlertType.CUSTOMER_ZERO and alert.affected_customers:
        cust_clean = sanitize_filename_part(alert.affected_customers[0])
        filename = f"{market_clean}_{channel_clean}_{cust_clean}_{date_str}.docx"
    else:
        filename = f"{market_clean}_{channel_clean}_{date_str}.docx"

    return os.path.join(base_dir, filename)


def set_cell_background(cell, hex_color: str):
    """Set shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa (1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_escalation_docx(
    alert: ConsolidatedAlert,
    investigation: DropInvestigationResult,
    output_dir: str = "escalations",
) -> str:
    """Generate a clean half-page Word document (.docx) matching the escalation email."""
    os.makedirs(output_dir, exist_ok=True)
    file_path = get_escalation_docx_path(alert, base_dir=output_dir)

    doc = Document()

    # Set tight margins so document fits cleanly on ~half page
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    esc_label = f"Day {alert.escalation_day}"
    if not alert.escalation_confirmed:
        esc_label += " (Unconfirmed)"

    # 1. Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    sub_run = title_p.add_run("ATLIQ HARDWARE • INCIDENT ESCALATION REPORT\n")
    sub_run.font.size = Pt(8.5)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGB_GRAY

    h1_run = title_p.add_run(f"{alert.market} {alert.channel} Sales Drop • ")
    h1_run.font.size = Pt(15)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGB_BLACK

    esc_run = title_p.add_run(esc_label)
    esc_run.font.size = Pt(15)
    esc_run.font.bold = True
    esc_run.font.color.rgb = RGB_RED

    # 2. Metadata Table
    copy_str = f"{alert.copy_owner.name} ({alert.copy_owner.role})" if alert.copy_owner else "None (Day 1 / Region Head)"
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Order Date:", str(alert.date)),
        ("Primary Owner:", f"{alert.primary_owner.name} ({alert.primary_owner.role})"),
        ("Escalation Copy:", copy_str),
    ]

    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(5.0)

        set_cell_background(c0, "F8F9FA")
        set_cell_background(c1, "F8F9FA")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGB_GRAY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9)
        r1.font.bold = True
        r1.font.color.rgb = RGB_BLACK

    # 3. Metrics Section
    if alert.market_channel_baseline is not None:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(2)
        r_sec = p_sec.add_run("PERFORMANCE VS 4-WEEK MEDIAN")
        r_sec.font.size = Pt(9.5)
        r_sec.font.bold = True
        r_sec.font.color.rgb = RGB_BLUE

        table = doc.add_table(rows=3, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        rows_data = [
            ("4-Week Median Baseline", f"${alert.market_channel_baseline:,.2f}", False, RGB_BLACK),
            ("Observed Day Sales", f"${alert.market_channel_actual:,.2f}", False, RGB_BLACK),
            ("Net Sales Decline", f"-{alert.market_channel_decline_pct}% [Needs attention]", True, RGB_RED),
        ]

        for i, (lbl, val, is_alert_row, color) in enumerate(rows_data):
            row = table.rows[i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(4.5)
            c1.width = Inches(2.3)

            bg_hex = "FFEBEE" if is_alert_row else ("F8F9FA" if i % 2 == 0 else "FFFFFF")
            set_cell_background(c0, bg_hex)
            set_cell_background(c1, bg_hex)
            set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
            set_cell_margins(c1, top=60, bottom=60, left=100, right=100)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(lbl)
            r0.font.size = Pt(9)
            if is_alert_row:
                r0.font.bold = True
                r0.font.color.rgb = RGB_RED

            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(val)
            r1.font.size = Pt(9)
            r1.font.bold = True
            r1.font.color.rgb = color

    # 4. Zero-Order Customers Section
    if alert.affected_customers:
        p_cust = doc.add_paragraph()
        p_cust.paragraph_format.space_before = Pt(8)
        p_cust.paragraph_format.space_after = Pt(2)
        r_cust = p_cust.add_run("CUSTOMER(S) STOPPED ORDERING")
        r_cust.font.size = Pt(9.5)
        r_cust.font.bold = True
        r_cust.font.color.rgb = RGB_RED

        for c_name in alert.affected_customers:
            base = alert.customer_baselines.get(c_name, 0.0)
            p_item = doc.add_paragraph(style="List Bullet")
            p_item.paragraph_format.space_before = Pt(1)
            p_item.paragraph_format.space_after = Pt(1)

            r_name = p_item.add_run(f"{c_name}: ")
            r_name.font.size = Pt(9)
            r_name.font.bold = True

            r_desc = p_item.add_run(f"Normal weekday baseline ${base:,.2f} -> ")
            r_desc.font.size = Pt(9)

            r_act = p_item.add_run("$0.00 actual [Stopped ordering]")
            r_act.font.size = Pt(9)
            r_act.font.bold = True
            r_act.font.color.rgb = RGB_RED

    # 5. Cause Analysis Section
    p_cause = doc.add_paragraph()
    p_cause.paragraph_format.space_before = Pt(8)
    p_cause.paragraph_format.space_after = Pt(2)
    r_cause = p_cause.add_run("CAUSE ANALYSIS & EVIDENCE")
    r_cause.font.size = Pt(9.5)
    r_cause.font.bold = True
    r_cause.font.color.rgb = RGB_BLUE

    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(2)
    p_head.paragraph_format.space_after = Pt(3)
    r_hl = p_head.add_run(investigation.headline)
    r_hl.font.size = Pt(9.5)
    r_hl.font.bold = True

    p_body = doc.add_paragraph()
    p_body.paragraph_format.space_before = Pt(0)
    p_body.paragraph_format.space_after = Pt(3)
    r_body = p_body.add_run(investigation.likely_causes)
    r_body.font.size = Pt(9)

    p_ev = doc.add_paragraph()
    p_ev.paragraph_format.space_before = Pt(0)
    p_ev.paragraph_format.space_after = Pt(3)
    r_ev_lbl = p_ev.add_run("Evidence: ")
    r_ev_lbl.font.size = Pt(8.5)
    r_ev_lbl.font.bold = True
    r_ev = p_ev.add_run(investigation.evidence)
    r_ev.font.size = Pt(8.5)
    r_ev.font.color.rgb = RGB_GRAY

    # Status footer note
    p_stat = doc.add_paragraph()
    p_stat.paragraph_format.space_before = Pt(4)
    p_stat.paragraph_format.space_after = Pt(0)
    r_stat = p_stat.add_run(f"Status: {investigation.analysis_status} ({investigation.confidence} confidence)")
    r_stat.font.size = Pt(8)
    r_stat.font.italic = True
    r_stat.font.color.rgb = RGB_GRAY

    doc.save(file_path)
    return file_path

