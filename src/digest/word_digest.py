"""One-page Weekly Digest Word document generator (.docx)."""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from src.context.stakeholders import Stakeholder
from src.digest.scorecard import WeeklyScorecard
from src.ai.digest_investigator import DigestNarrativeResult
from src.reports.formatting import (
    RGB_RED,
    RGB_GREEN,
    RGB_BLUE,
    RGB_GRAY,
    RGB_BLACK,
    format_delta_text,
)


def sanitize_filename_part(part: str) -> str:
    """Sanitize string for Windows filesystem safety."""
    cleaned = re.sub(r'[\\/*?:"<>|]', "", part)
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned.strip("._")


def get_digest_docx_path(stakeholder: Stakeholder, scorecard: WeeklyScorecard, base_dir: str = "digests") -> str:
    """Generate sanitized output path for digest .docx file.
    
    Format: digests/<name>_<monday>_to_<sunday>.docx
    """
    clean_name = sanitize_filename_part(stakeholder.name)
    m_str = scorecard.monday.isoformat()
    s_str = scorecard.sunday.isoformat()
    filename = f"{clean_name}_{m_str}_to_{s_str}.docx"
    return os.path.join(base_dir, filename)


def set_cell_background(cell, hex_color: str):
    """Set shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Set cell padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_digest_docx(
    stakeholder: Stakeholder,
    scorecard: WeeklyScorecard,
    narrative: DigestNarrativeResult,
    output_dir: str = "digests",
) -> str:
    """Generate a clean one-page Word document (.docx) adhering to PRD 7.4 delivery order."""
    os.makedirs(output_dir, exist_ok=True)
    file_path = get_digest_docx_path(stakeholder, scorecard, base_dir=output_dir)

    doc = Document()

    # Tight margins to ensure 1-page fit
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    # 1. Header Banner
    head_p = doc.add_paragraph()
    head_p.paragraph_format.space_before = Pt(0)
    head_p.paragraph_format.space_after = Pt(2)
    sub = head_p.add_run("ATLIQ HARDWARE • COMPANY-WIDE WEEKLY EXECUTIVE DIGEST\n")
    sub.font.size = Pt(8.5)
    sub.font.bold = True
    sub.font.color.rgb = RGB_GRAY

    main_title = head_p.add_run(f"Weekly Brief for {stakeholder.name} ({stakeholder.role})")
    main_title.font.size = Pt(14)
    main_title.font.bold = True
    main_title.font.color.rgb = RGB_BLACK

    date_sub = head_p.add_run(f" • Week: {scorecard.monday} to {scorecard.sunday}")
    date_sub.font.size = Pt(10)
    date_sub.font.color.rgb = RGB_BLUE

    # 2. Section 1: Scorecard Table
    sc_title = doc.add_paragraph()
    sc_title.paragraph_format.space_before = Pt(4)
    sc_title.paragraph_format.space_after = Pt(2)
    r_sc = sc_title.add_run("1. COMPANY SCORECARD")
    r_sc.font.size = Pt(9.5)
    r_sc.font.bold = True
    r_sc.font.color.rgb = RGB_BLUE

    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.0), Inches(1.8), Inches(1.8), Inches(1.6)]
    headers = ["Metric", "This Week", "vs Prior Week", "vs Target / Impact"]

    # Header Row
    hdr_row = table.rows[0]
    for idx, name in enumerate(headers):
        c = hdr_row.cells[idx]
        c.width = col_widths[idx]
        set_cell_background(c, "F0F4F8")
        set_cell_margins(c, top=50, bottom=50, left=80, right=80)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(name)
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGB_BLACK

    # Rows Data
    # Net Sales: compared with target
    ns_lbl, ns_hex, ns_rgb = format_delta_text(scorecard.net_sales_change_pct, "%")
    tar_lbl, tar_hex, tar_rgb = format_delta_text(scorecard.net_sales_vs_target_pct, "%")

    # Gross Margin %
    gm_lbl, gm_hex, gm_rgb = format_delta_text(scorecard.gross_margin_pct_change, " pts")

    # Units
    unit_lbl, unit_hex, unit_rgb = format_delta_text(scorecard.units_change_pct, "%")

    # Gross Margin $
    money_lbl, money_hex, money_rgb = format_delta_text(scorecard.money_impact, "$")

    data_rows = [
        ("Net Sales", f"${scorecard.net_sales:,.2f}", ns_lbl, ns_rgb, f"{tar_lbl} (Target: ${scorecard.net_sales_target:,.2f})", tar_rgb),
        ("Gross Margin %", f"{scorecard.gross_margin_pct:.1f}%", gm_lbl, gm_rgb, "— (No target)", RGB_GRAY),
        ("Units Sold", f"{scorecard.units_sold:,}", unit_lbl, unit_rgb, "— (No target)", RGB_GRAY),
        ("Gross Margin $", f"${scorecard.gross_margin:,.2f}", money_lbl, money_rgb, f"Monthly: ${scorecard.monthly_projected_impact:,.2f}", money_rgb),
    ]

    for row_idx, (m_name, val_str, wow_str, wow_rgb, target_str, tar_color) in enumerate(data_rows, 1):
        row = table.rows[row_idx]
        bg = "FFFFFF" if row_idx % 2 == 1 else "FBFBFB"
        for c_idx in range(4):
            c = row.cells[c_idx]
            c.width = col_widths[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, top=45, bottom=45, left=80, right=80)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)

            if c_idx == 0:
                r = p.add_run(m_name)
                r.font.size = Pt(8.5)
                r.font.bold = True
            elif c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(val_str)
                r.font.size = Pt(8.5)
                r.font.bold = True
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(wow_str)
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = wow_rgb
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(target_str)
                r.font.size = Pt(8.0)
                r.font.color.rgb = tar_color

    # 3. Section 2: Headline
    hl_p = doc.add_paragraph()
    hl_p.paragraph_format.space_before = Pt(6)
    hl_p.paragraph_format.space_after = Pt(2)
    r_hl_lbl = hl_p.add_run("2. EXECUTIVE HEADLINE\n")
    r_hl_lbl.font.size = Pt(9)
    r_hl_lbl.font.bold = True
    r_hl_lbl.font.color.rgb = RGB_BLUE
    r_hl = hl_p.add_run(narrative.headline)
    r_hl.font.size = Pt(10)
    r_hl.font.bold = True

    # 4. Section 3: Insights
    ins_p = doc.add_paragraph()
    ins_p.paragraph_format.space_before = Pt(4)
    ins_p.paragraph_format.space_after = Pt(2)
    r_ins = ins_p.add_run("3. KEY INSIGHTS (PRIORITIZED FOR YOU)")
    r_ins.font.size = Pt(9)
    r_ins.font.bold = True
    r_ins.font.color.rgb = RGB_BLUE

    for item in narrative.insights[:3]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(item)
        r.font.size = Pt(8.5)

    # 5. Section 4: Causes
    cause_p = doc.add_paragraph()
    cause_p.paragraph_format.space_before = Pt(4)
    cause_p.paragraph_format.space_after = Pt(2)
    r_c_lbl = cause_p.add_run("4. ROOT CAUSES & EVIDENCE\n")
    r_c_lbl.font.size = Pt(9)
    r_c_lbl.font.bold = True
    r_c_lbl.font.color.rgb = RGB_BLUE
    r_c = cause_p.add_run(narrative.causes)
    r_c.font.size = Pt(8.5)

    # 6. Section 5: Actions
    act_p = doc.add_paragraph()
    act_p.paragraph_format.space_before = Pt(4)
    act_p.paragraph_format.space_after = Pt(2)
    r_act = act_p.add_run("5. RECOMMENDED ACTIONS")
    r_act.font.size = Pt(9)
    r_act.font.bold = True
    r_act.font.color.rgb = RGB_BLUE

    for act in narrative.actions[:2]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(act)
        r.font.size = Pt(8.5)

    # 7. Section 6: Decision Needed
    dec_p = doc.add_paragraph()
    dec_p.paragraph_format.space_before = Pt(4)
    dec_p.paragraph_format.space_after = Pt(2)
    r_dec_lbl = dec_p.add_run("6. DECISION NEEDED: ")
    r_dec_lbl.font.size = Pt(9)
    r_dec_lbl.font.bold = True
    r_dec_lbl.font.color.rgb = RGB_RED
    r_dec = dec_p.add_run(narrative.decision_needed)
    r_dec.font.size = Pt(8.5)
    r_dec.font.bold = True

    # 8. Section 7: Analysis Status / Caveat
    stat_p = doc.add_paragraph()
    stat_p.paragraph_format.space_before = Pt(4)
    stat_p.paragraph_format.space_after = Pt(0)
    r_st = stat_p.add_run(f"7. STATUS: {narrative.analysis_status}. Generated at 8:00 AM by AtliQ Morning Digest Agent.")
    r_st.font.size = Pt(7.5)
    r_st.font.italic = True
    r_st.font.color.rgb = RGB_GRAY

    doc.save(file_path)
    return file_path

